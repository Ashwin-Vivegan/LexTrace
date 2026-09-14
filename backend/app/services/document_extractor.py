import os
# pyrefly: ignore [missing-import]
import pypdf
# pyrefly: ignore [missing-import]
import docx
from typing import List, Dict, Any, Optional

class ExtractionError(Exception):
    """Raised when text extraction fails or when a file contains no extractable text."""
    pass

class DocumentExtractor:
    """
    Service responsible for format detection, file signature validation,
    and text extraction for PDF, DOCX, and TXT legal documents.
    Supports both flat string extraction and page-indexed structured block extraction.
    """

    @classmethod
    def validate_file_header(cls, file_path: str, file_type: str) -> None:
        """
        Inspects file magic bytes/header signatures to prevent file extension spoofing.
        """
        if not os.path.exists(file_path):
            raise ExtractionError(f"File not found: {file_path}")

        file_size = os.path.getsize(file_path)
        if file_size == 0:
            raise ExtractionError("Uploaded file is empty (0 bytes).")

        with open(file_path, "rb") as f:
            header = f.read(16)

        if file_type == "pdf":
            if not header.startswith(b"%PDF"):
                raise ExtractionError("Invalid PDF file format. Header magic bytes mismatch.")
        elif file_type == "docx":
            if not header.startswith(b"PK\x03\x04"):
                raise ExtractionError("Invalid DOCX file format. Header magic bytes mismatch.")

    @classmethod
    def extract_pdf_pages(cls, file_path: str) -> List[Dict[str, Any]]:
        """
        Extracts page-indexed text blocks from a PDF file using pypdf.
        Returns a list of dicts: [{"page_number": 1, "text": "..."}, ...]
        Raises ExtractionError if PDF is corrupted or contains zero extractable text.
        """
        try:
            reader = pypdf.PdfReader(file_path)
            page_blocks = []
            
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text and page_text.strip():
                    page_blocks.append({
                        "page_number": i + 1,
                        "text": page_text.strip()
                    })

            if not page_blocks:
                raise ExtractionError(
                    "PDF file contains no extractable text. It may be scanned or image-only (OCR is not enabled)."
                )
            return page_blocks
        except ExtractionError:
            raise
        except Exception as e:
            raise ExtractionError(f"Error parsing PDF file: {str(e)}")

    @classmethod
    def extract_pdf(cls, file_path: str) -> str:
        """Extracts full text from PDF as a single concatenated string."""
        pages = cls.extract_pdf_pages(file_path)
        return "\n\n".join(p["text"] for p in pages)

    @classmethod
    def extract_docx(cls, file_path: str) -> str:
        """
        Extracts text from paragraphs and tables in a DOCX file using python-docx.
        """
        try:
            doc = docx.Document(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            
            # Extract cell texts from any embedded tables
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_content = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_content:
                        table_texts.append(" | ".join(row_content))

            full_content = []
            if paragraphs:
                full_content.extend(paragraphs)
            if table_texts:
                full_content.append("--- Document Tables ---")
                full_content.extend(table_texts)

            full_text = "\n\n".join(full_content).strip()
            if not full_text:
                raise ExtractionError("DOCX file contains no extractable text.")
            return full_text
        except ExtractionError:
            raise
        except Exception as e:
            raise ExtractionError(f"Error parsing DOCX file: {str(e)}")

    @classmethod
    def extract_txt(cls, file_path: str) -> str:
        """
        Extracts text from a plain text (TXT) file with UTF-8 / Latin-1 fallback decoding.
        """
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read().strip()
        except UnicodeDecodeError:
            try:
                with open(file_path, "r", encoding="latin-1") as f:
                    text = f.read().strip()
            except Exception as e:
                raise ExtractionError(f"Failed to decode text file: {str(e)}")

        if not text:
            raise ExtractionError("TXT file is empty or contains no readable text.")
        return text

    @classmethod
    def extract_text(cls, file_path: str, file_type: str) -> str:
        """
        Main entry point for flat string text extraction.
        """
        file_type_lower = file_type.lower().strip()
        cls.validate_file_header(file_path, file_type_lower)

        if file_type_lower == "pdf":
            return cls.extract_pdf(file_path)
        elif file_type_lower == "docx":
            return cls.extract_docx(file_path)
        elif file_type_lower == "txt":
            return cls.extract_txt(file_path)
        else:
            raise ExtractionError(f"Unsupported file format: .{file_type_lower}")

    @classmethod
    def extract_structured_pages(cls, file_path: str, file_type: str) -> List[Dict[str, Any]]:
        """
        Extracts structured page blocks from file.
        For PDF: returns list of [{"page_number": 1, "text": "..."}, ...]
        For DOCX/TXT: returns [{"page_number": None, "text": full_text}]
        """
        file_type_lower = file_type.lower().strip()
        cls.validate_file_header(file_path, file_type_lower)

        if file_type_lower == "pdf":
            return cls.extract_pdf_pages(file_path)
        else:
            full_text = cls.extract_text(file_path, file_type_lower)
            return [{"page_number": None, "text": full_text}]
